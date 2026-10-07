"""
Script to generate multi-format mock documents for Company Y.
Formats: PDF (.pdf), Word (.docx), Plain Text (.txt), Markdown (.md).
"""

from pathlib import Path
from fpdf import FPDF
import docx

PROJECT_ROOT = Path(__file__).resolve().parent.parent
COMPANY_Y_DIR = PROJECT_ROOT / "data" / "company_y"
COMPANY_Y_UPDATED_DIR = COMPANY_Y_DIR / "updated"

COMPANY_Y_DIR.mkdir(parents=True, exist_ok=True)
COMPANY_Y_UPDATED_DIR.mkdir(parents=True, exist_ok=True)


def create_pdf(file_path: Path, title: str, sections: list):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("helvetica", "B", 14)
    pdf.cell(0, 10, title, ln=True, align="L")
    pdf.ln(4)
    for heading, body in sections:
        pdf.set_font("helvetica", "B", 11)
        pdf.cell(0, 7, heading, ln=True)
        pdf.set_font("helvetica", "", 9.5)
        pdf.multi_cell(0, 5, body)
        pdf.ln(3)
    pdf.output(str(file_path))


def create_docx(file_path: Path, title: str, sections: list):
    doc = docx.Document()
    doc.add_heading(title, level=1)
    for heading, body in sections:
        doc.add_heading(heading, level=2)
        doc.add_paragraph(body)
    doc.save(str(file_path))


def main():
    # 1. Baseline Data Boundary Policy (PDF) - Intentional MISMATCH: Global Cloud / Non-US allowed
    data_boundary_v1 = [
        ("1. Purpose and Scope", "This policy defines cloud hosting and data boundary controls for Company Y AI models and telemetry."),
        ("2. Cloud Infrastructure Architecture", "Company Y hosts production AI services across commercial multi-cloud infrastructure, including AWS us-east-1 and global EU regions (eu-central-1, eu-west-1) to ensure global latency optimization and active-active multi-region resiliency."),
        ("3. Cross-Border Transit & Redundancy", "In case of primary US regional outages, automated cloud DNS automatically routes API requests and model training pipelines to overseas international nodes. Data residency boundaries are flexible and subject to global availability requirements."),
    ]
    create_pdf(COMPANY_Y_DIR / "data_boundary_policy.pdf", "Company Y - Cloud Infrastructure & Data Boundary Policy v1.0", data_boundary_v1)

    # 2. Baseline Access Control Policy (DOCX) - Intentional PARTIAL: Password + SMS MFA, no US-person citizenship constraint
    access_control_v1 = [
        ("1. Identity and Access Governance", "Company Y enforces role-based access control (RBAC) across corporate and AI platform accounts."),
        ("2. Authentication Mechanism", "Standard user authentication requires corporate email passwords supplemented by SMS or time-based OTP (TOTP) application tokens."),
        ("3. Personnel Screening & Global Contractors", "Access to AI training datasets and model APIs is granted to verified full-time employees and international offshore software engineering contractors under NDA. US citizenship is not formally verified or mandated."),
        ("4. Privilege Review Cycle", "Access entitlements are reviewed by engineering managers on an annual (12-month) recurring schedule."),
    ]
    create_docx(COMPANY_Y_DIR / "access_control_policy.docx", "Company Y - Access Control & Authentication Policy v1.0", access_control_v1)

    # 3. Baseline Audit Logging Policy (TXT) - Intentional PARTIAL: 90 days mutable retention
    audit_logging_v1_text = """# Company Y - Audit Logging and Monitoring Policy v1.0

1. Overview:
Company Y records operational event logs from API servers, databases, and container orchestration clusters.

2. Retention Schedule:
System and audit logs are retained in cloud object storage for a rolling duration of 90 days. Logs older than 90 days are automatically pruned to manage infrastructure storage costs.

3. Log Storage and Tamper Resistance:
Logs are stored in standard standard-tier cloud buckets. Bucket administrators possess permissions to archive, rotate, or delete log buckets during routine system maintenance. WORM or immutable lock modes are not currently implemented.
"""
    (COMPANY_Y_DIR / "audit_logging_policy.txt").write_text(audit_logging_v1_text, encoding="utf-8")

    # 4. Baseline Encryption Policy (MD) - MATCH
    encryption_v1_text = """# Company Y - Cryptographic and Encryption Policy v1.0

## 1. Cryptographic Standards
All Company Y data environments must enforce robust cryptographic protections.

## 2. Encryption in Transit
All network transmissions between clients and AI model endpoints must enforce TLS 1.3 encryption utilizing NIST/FIPS 140-2 validated cryptographic modules and cipher suites.

## 3. Encryption at Rest
All persistent storage volumes, model weight repositories, and inference caches must be encrypted at rest using AES-256 bit encryption keys managed through customer KMS with annual key rotation.
"""
    (COMPANY_Y_DIR / "encryption_policy.md").write_text(encryption_v1_text, encoding="utf-8")

    # 5. Baseline Incident Response Policy (MD) - MATCH
    ir_v1_text = """# Company Y - Incident Response Playbook v1.0

## 1. Incident Classification
Security incidents involving unauthorized access, data compromise, or infrastructure breaches are classified as Severity-1 Critical.

## 2. Mandatory Reporting Window
In compliance with defense workloads and partner obligations, all confirmed security breaches involving sensitive data or AI environments must be formally reported to affected defense authorities and partner incident response teams within 72 hours of discovery.

## 3. Forensic Preservation
Digital forensic logs, container memory dumps, and audit artifacts must be quarantined and preserved in tamper-evident storage for at least 90 days following incident resolution.
"""
    (COMPANY_Y_DIR / "incident_response_policy.md").write_text(ir_v1_text, encoding="utf-8")

    # 6. Baseline Vulnerability Management Policy (MD) - MATCH
    vuln_v1_text = """# Company Y - Vulnerability Management Policy v1.0

## 1. Continuous Vulnerability Scanning
All container images, virtual machines, and AI code libraries must undergo automated vulnerability scanning in continuous integration prior to deployment.

## 2. Remediation Timelines
Vulnerabilities discovered in production AI workloads must be remediated according to severity:
- Critical vulnerabilities (CVSS >= 9.0) must be patched or mitigated within 14 calendar days.
- High vulnerabilities (CVSS 7.0 - 8.9) must be remediated within 30 calendar days.
"""
    (COMPANY_Y_DIR / "vulnerability_management_policy.md").write_text(vuln_v1_text, encoding="utf-8")

    # 7. Baseline CUI Handling Policy (MD) - MATCH
    cui_v1_text = """# Company Y - Controlled Unclassified Information (CUI) Policy v1.0

## 1. CUI Identification & Markings
All digital data containing Controlled Unclassified Information (CUI) processed within AI environments must bear clear classification banners and metadata tags.

## 2. Media Sanitization
Any physical storage or decommissioned virtual disk instances containing CUI must be cryptographically wiped in strict conformance with NIST SP 800-88 Rev. 1 Guidelines for Media Sanitization prior to deallocation.
"""
    (COMPANY_Y_DIR / "cui_handling_policy.md").write_text(cui_v1_text, encoding="utf-8")

    # =========================================================================
    # REMEDIATED POLICIES (v2) for Company Y
    # =========================================================================

    # Remediated Data Boundary Policy (PDF) - MATCH: FedRAMP High / AWS GovCloud US-only, strictly CONUS
    data_boundary_v2 = [
        ("1. Purpose and Scope", "This remediated policy defines strict US-only sovereign boundary controls for Company Y defense AI infrastructure."),
        ("2. Sovereign US Data Boundary & GovCloud Hosting", "All customer training data, foundation models, fine-tuned weights, and AI inference payloads must reside and execute strictly within FedRAMP High or AWS GovCloud (US) regions located in the continental United States (CONUS)."),
        ("3. Strict Prohibition of Cross-Border Transit", "Cross-border data transit, replication, backup, or routing outside CONUS boundaries is strictly prohibited. International cloud regions, foreign IP routing, and offshore fallback mechanisms are completely decommissioned."),
    ]
    create_pdf(COMPANY_Y_UPDATED_DIR / "data_boundary_policy_v2.pdf", "Company Y - Sovereign US Data Boundary Policy v2.0 (Remediated)", data_boundary_v2)

    # Remediated Access Control Policy (DOCX) - MATCH: US-person verification, FIDO2 MFA, 90-day review
    access_control_v2 = [
        ("1. Strict US-Person Access Control", "Access to Company Y defense AI systems, training data, and infrastructure is strictly restricted to verified U.S. persons (U.S. citizens or lawful permanent residents). Proof of citizenship or permanent residency is verified by HR compliance prior to provisioning."),
        ("2. Phishing-Resistant Hardware MFA", "All administrative and privileged access to AI workloads mandates phishing-resistant hardware tokens (FIDO2 / WebAuthn). SMS and basic TOTP authenticators are strictly prohibited for privileged roles."),
        ("3. Least Privilege and Quarterly Review", "Access permissions operate under least privilege role-based access control (RBAC). All user entitlements and access grants must undergo mandatory formal re-certification every 90 days."),
    ]
    create_docx(COMPANY_Y_UPDATED_DIR / "access_control_policy_v2.docx", "Company Y - Access Control and Authentication Policy v2.0 (Remediated)", access_control_v2)

    # Remediated Audit Logging Policy (TXT) - MATCH: 365-day retention in immutable WORM storage
    audit_logging_v2_text = """# Company Y - Audit Logging and Monitoring Policy v2.0 (Remediated)

1. Purpose:
Establishes compliant, immutable audit logging for defense AI workloads.

2. Immutable Storage (WORM):
All audit events, authentication attempts, API calls, and administrative modifications must be written to append-only immutable storage using AWS S3 Object Lock in Compliance Mode (Write-Once-Read-Many / WORM). Modification, truncation, or deletion of log records by any user or administrator is cryptographically prevented.

3. Mandatory 365-Day Retention Schedule:
All audit and security telemetry logs must be preserved for a minimum retention duration of 365 days (1 year). Real-time log forwarding to the centralized Security Information and Event Management (SIEM) system is enforced with synchronized NTP timestamps.
"""
    (COMPANY_Y_UPDATED_DIR / "audit_logging_policy_v2.txt").write_text(audit_logging_v2_text, encoding="utf-8")

    print("Successfully generated multi-format mock policies for Company Y (v1 and v2)!")


if __name__ == "__main__":
    main()
