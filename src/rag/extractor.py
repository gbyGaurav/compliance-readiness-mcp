"""
Multi-format document text extraction engine.
Supports PDF (.pdf), Word Documents (.docx), Plain Text (.txt), and Markdown (.md).
"""

from pathlib import Path
from typing import Optional


def extract_text_from_pdf(file_path: Path) -> str:
    """Extracts text from a PDF file using pypdf."""
    import pypdf
    reader = pypdf.PdfReader(str(file_path))
    pages_text = []
    for i, page in enumerate(reader.pages):
        page_content = page.extract_text() or ""
        if page_content.strip():
            pages_text.append(f"## Page {i + 1}\n\n{page_content.strip()}")
    return "\n\n".join(pages_text)


def extract_text_from_docx(file_path: Path) -> str:
    """Extracts text from a DOCX file using python-docx."""
    import docx
    doc = docx.Document(str(file_path))
    paragraphs_text = []
    for p in doc.paragraphs:
        text = p.text.strip()
        if text:
            # Preserve headings if styled as heading
            if p.style and p.style.name.startswith("Heading"):
                level = p.style.name.replace("Heading", "").strip() or "2"
                paragraphs_text.append(f"{'#' * int(level if level.isdigit() else 2)} {text}")
            else:
                paragraphs_text.append(text)
    
    # Also extract tables if present
    for table in doc.tables:
        table_rows = []
        for row in table.rows:
            row_cells = [cell.text.strip() for cell in row.cells]
            table_rows.append(" | ".join(row_cells))
        if table_rows:
            paragraphs_text.append("\n".join(table_rows))

    return "\n\n".join(paragraphs_text)


def extract_text_from_file(file_path: Path) -> str:
    """
    Extracts plain text from supported document types (.pdf, .docx, .txt, .md).
    Raises ValueError for unsupported file formats.
    """
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"Document file not found: {path}")

    ext = path.suffix.lower()
    if ext == ".pdf":
        return extract_text_from_pdf(path)
    elif ext == ".docx":
        return extract_text_from_docx(path)
    elif ext in [".txt", ".md", ".markdown"]:
        return path.read_text(encoding="utf-8", errors="replace")
    else:
        raise ValueError(f"Unsupported document format '{ext}'. Supported formats: .pdf, .docx, .txt, .md")
